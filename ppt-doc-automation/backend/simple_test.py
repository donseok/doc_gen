"""간단한 템플릿 분석 및 생성 테스트"""
import os
import sys

# 현재 디렉토리를 확인하고 출력
print(f"Current directory: {os.getcwd()}")
print(f"Script directory: {os.path.dirname(os.path.abspath(__file__))}")

# pptx 라이브러리 임포트
try:
    from pptx import Presentation
    from pptx.util import Pt
    print("pptx library imported successfully")
except ImportError as e:
    print(f"Failed to import pptx: {e}")
    sys.exit(1)

# 템플릿 경로
TEMPLATE_PATH = r"d:\doc_gen\ppt-doc-automation\templates\pptx\group\그룹 공통 PPT템플릿.pptx"
OUTPUT_PATH = r"d:\doc_gen\ppt-doc-automation\backend\output\simple_test_output.pptx"

print(f"\nTemplate exists: {os.path.exists(TEMPLATE_PATH)}")

# 템플릿 열기
print("\nOpening template...")
prs = Presentation(TEMPLATE_PATH)

print(f"Slide count: {len(prs.slides)}")
print(f"Layout count: {len(prs.slide_layouts)}")

# 레이아웃 목록 출력
print("\nLayouts:")
for i, layout in enumerate(prs.slide_layouts):
    print(f"  [{i}] {layout.name}")

# 새 슬라이드 추가 테스트 (첫 번째 레이아웃 - 표지)
print("\nAdding test slide with layout 0 (Title)...")
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)

# 플레이스홀더 분석
print("\nPlaceholders in new slide:")
for shape in slide.shapes:
    info = f"  - {shape.shape_id}: {shape.name}"
    if hasattr(shape, 'placeholder_format') and shape.placeholder_format:
        info += f" (idx={shape.placeholder_format.idx})"
    if shape.has_text_frame:
        info += " [has text]"
    print(info)

# 저장
print(f"\nSaving to: {OUTPUT_PATH}")
os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
prs.save(OUTPUT_PATH)
print("Done!")
