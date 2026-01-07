"""
카드 컴포넌트

PPT 슬라이드에 다양한 카드 요소를 생성합니다.

카드 유형:
    - BigNumberCard: 대형 숫자 표시 (KPI)
    - IconCard: 아이콘 + 제목 + 설명
    - FeatureCard: 번호 + 제목 + 설명
    - CardGrid: 카드 그리드 레이아웃
"""

from typing import Dict, Any, List, Optional
from pptx.slide import Slide
from pptx.util import Emu, Pt, Cm, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from ..base import BaseComponent, RenderContext
from ..factory import register_component


@register_component('big_numbers', category='visual')
class BigNumberCardComponent(BaseComponent):
    """대형 숫자 카드 컴포넌트
    
    KPI나 주요 지표를 큰 숫자로 표시합니다.
    
    Data Schema:
        {
            "items": [
                {
                    "value": "500만",
                    "label": "연간 물류 처리량",
                    "suffix": "건",  # 선택적
                    "change": "+30%"  # 선택적
                },
                ...
            ],
            "columns": 3  # 열 수, 선택적
        }
    """
    
    CARD_HEIGHT = Emu(Cm(3.5))
    CARD_MARGIN = Emu(Cm(0.3))
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any],
        top: Emu,
        **kwargs
    ) -> Emu:
        """대형 숫자 카드 렌더링
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            data: 카드 데이터
            top: 시작 Y 위치
            
        Returns:
            사용된 높이 (EMU)
        """
        items = data.get('items', [])
        columns = data.get('columns', min(len(items), 4))
        
        if not items:
            return Emu(0)
        
        # 카드 크기 계산
        total_margin = int(self.CARD_MARGIN) * (columns - 1)
        card_width = Emu((int(context.content_width) - total_margin) // columns)
        
        left_start = context.get_content_start_x()
        
        for i, item in enumerate(items):
            col = i % columns
            row = i // columns
            
            card_left = Emu(int(left_start) + col * (int(card_width) + int(self.CARD_MARGIN)))
            card_top = Emu(int(top) + row * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
            
            self._render_single_card(slide, context, item, card_left, card_top, card_width)
        
        # 총 행 수 계산
        num_rows = (len(items) + columns - 1) // columns
        total_height = Emu(num_rows * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
        
        return total_height
    
    def get_required_height(
        self, 
        context: RenderContext, 
        data: Dict[str, Any],
        **kwargs
    ) -> Emu:
        """필요한 높이 계산"""
        items = data.get('items', [])
        columns = data.get('columns', min(len(items), 4))
        
        if not items:
            return Emu(0)
        
        num_rows = (len(items) + columns - 1) // columns
        return Emu(num_rows * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
    
    def _render_single_card(
        self,
        slide: Slide,
        context: RenderContext,
        item: Dict[str, Any],
        left: Emu,
        top: Emu,
        width: Emu
    ) -> None:
        """단일 대형 숫자 카드 렌더링
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            item: 카드 아이템 데이터
            left, top, width: 위치 및 너비
        """
        # 카드 배경
        card_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, self.CARD_HEIGHT
        )
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = context.theme.colors.background
        card_bg.line.color.rgb = context.theme.colors.secondary
        card_bg.line.width = Pt(1)
        
        # 큰 숫자
        value = item.get('value', '')
        suffix = item.get('suffix', '')
        value_text = f"{value}{suffix}"
        
        value_top = Emu(int(top) + Emu(Cm(0.5)))
        value_height = Emu(Cm(1.5))
        
        value_shape = slide.shapes.add_textbox(left, value_top, width, value_height)
        tf = value_shape.text_frame
        p = tf.paragraphs[0]
        p.text = value_text
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.title_font
        run.font.size = Pt(36)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.primary
        
        # 라벨
        label = item.get('label', '')
        label_top = Emu(int(top) + Emu(Cm(2.2)))
        label_height = Emu(Cm(0.8))
        
        label_shape = slide.shapes.add_textbox(left, label_top, width, label_height)
        tf = label_shape.text_frame
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(11)
        run.font.color.rgb = context.theme.colors.text_secondary
        
        # 변화량 (선택적)
        change = item.get('change', '')
        if change:
            change_top = Emu(int(top) + Emu(Cm(2.9)))
            change_height = Emu(Cm(0.5))
            
            change_shape = slide.shapes.add_textbox(left, change_top, width, change_height)
            tf = change_shape.text_frame
            p = tf.paragraphs[0]
            p.text = change
            p.alignment = PP_ALIGN.CENTER
            
            run = p.runs[0] if p.runs else p.add_run()
            run.font.name = context.theme.fonts.body_font
            run.font.size = Pt(10)
            
            # 증가/감소에 따른 색상
            if change.startswith('+'):
                run.font.color.rgb = RGBColor(0, 150, 0)  # 초록
            elif change.startswith('-'):
                run.font.color.rgb = context.theme.colors.accent  # 빨강


@register_component('icon_cards', category='visual')
class IconCardComponent(BaseComponent):
    """아이콘 카드 컴포넌트
    
    아이콘(이모지) + 제목 + 설명 형태의 카드입니다.
    
    Data Schema:
        {
            "items": [
                {
                    "icon": "🚀",  # 이모지 또는 아이콘 이름
                    "title": "빠른 처리",
                    "description": "실시간 배송 추적 지원"
                },
                ...
            ],
            "columns": 3
        }
    """
    
    CARD_HEIGHT = Emu(Cm(4))
    CARD_MARGIN = Emu(Cm(0.3))
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any],
        top: Emu,
        **kwargs
    ) -> Emu:
        """아이콘 카드 렌더링"""
        items = data.get('items', [])
        columns = data.get('columns', min(len(items), 4))
        
        if not items:
            return Emu(0)
        
        total_margin = int(self.CARD_MARGIN) * (columns - 1)
        card_width = Emu((int(context.content_width) - total_margin) // columns)
        left_start = context.get_content_start_x()
        
        for i, item in enumerate(items):
            col = i % columns
            row = i // columns
            
            card_left = Emu(int(left_start) + col * (int(card_width) + int(self.CARD_MARGIN)))
            card_top = Emu(int(top) + row * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
            
            self._render_single_card(slide, context, item, card_left, card_top, card_width)
        
        num_rows = (len(items) + columns - 1) // columns
        return Emu(num_rows * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
    
    def get_required_height(
        self, 
        context: RenderContext, 
        data: Dict[str, Any],
        **kwargs
    ) -> Emu:
        """필요한 높이 계산"""
        items = data.get('items', [])
        columns = data.get('columns', min(len(items), 4))
        
        if not items:
            return Emu(0)
        
        num_rows = (len(items) + columns - 1) // columns
        return Emu(num_rows * (int(self.CARD_HEIGHT) + int(self.CARD_MARGIN)))
    
    def _render_single_card(
        self,
        slide: Slide,
        context: RenderContext,
        item: Dict[str, Any],
        left: Emu,
        top: Emu,
        width: Emu
    ) -> None:
        """단일 아이콘 카드 렌더링"""
        # 카드 배경
        card_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, self.CARD_HEIGHT
        )
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = context.theme.colors.background
        card_bg.line.color.rgb = context.theme.colors.secondary
        card_bg.line.width = Pt(1)
        
        # 아이콘 (이모지)
        icon = item.get('icon', '📌')
        icon_top = Emu(int(top) + Emu(Cm(0.3)))
        icon_height = Emu(Cm(1.2))
        
        icon_shape = slide.shapes.add_textbox(left, icon_top, width, icon_height)
        tf = icon_shape.text_frame
        p = tf.paragraphs[0]
        p.text = icon
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(28)
        
        # 제목
        title = item.get('title', '')
        title_top = Emu(int(top) + Emu(Cm(1.6)))
        title_height = Emu(Cm(0.8))
        
        title_shape = slide.shapes.add_textbox(left, title_top, width, title_height)
        tf = title_shape.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.primary
        
        # 설명
        description = item.get('description', '')
        if description:
            desc_top = Emu(int(top) + Emu(Cm(2.5)))
            desc_height = Emu(Cm(1.3))
            
            desc_shape = slide.shapes.add_textbox(
                Emu(int(left) + Emu(Cm(0.2))),
                desc_top,
                Emu(int(width) - Emu(Cm(0.4))),
                desc_height
            )
            tf = desc_shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = description
            p.alignment = PP_ALIGN.CENTER
            
            run = p.runs[0] if p.runs else p.add_run()
            run.font.name = context.theme.fonts.caption_font
            run.font.size = Pt(10)
            run.font.color.rgb = context.theme.colors.text_secondary


@register_component('feature_cards', category='visual')
class FeatureCardComponent(BaseComponent):
    """기능 카드 컴포넌트
    
    번호 + 제목 + 설명 형태의 카드입니다.
    주로 핵심 기능이나 단계를 표시할 때 사용합니다.
    
    Data Schema:
        {
            "items": [
                {
                    "number": "01",
                    "title": "입고 관리",
                    "description": "물류 입고 프로세스 자동화"
                },
                ...
            ],
            "columns": 2,
            "style": "horizontal"  # horizontal 또는 vertical
        }
    """
    
    CARD_HEIGHT_HORIZONTAL = Emu(Cm(2.5))
    CARD_HEIGHT_VERTICAL = Emu(Cm(4))
    CARD_MARGIN = Emu(Cm(0.3))
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any],
        top: Emu,
        **kwargs
    ) -> Emu:
        """기능 카드 렌더링"""
        items = data.get('items', [])
        columns = data.get('columns', 2)
        style = data.get('style', 'horizontal')
        
        card_height = self.CARD_HEIGHT_VERTICAL if style == 'vertical' else self.CARD_HEIGHT_HORIZONTAL
        
        if not items:
            return Emu(0)
        
        total_margin = int(self.CARD_MARGIN) * (columns - 1)
        card_width = Emu((int(context.content_width) - total_margin) // columns)
        left_start = context.get_content_start_x()
        
        for i, item in enumerate(items):
            col = i % columns
            row = i // columns
            
            card_left = Emu(int(left_start) + col * (int(card_width) + int(self.CARD_MARGIN)))
            card_top = Emu(int(top) + row * (int(card_height) + int(self.CARD_MARGIN)))
            
            # 기본값으로 번호 설정
            if 'number' not in item:
                item['number'] = f"{i+1:02d}"
            
            if style == 'vertical':
                self._render_vertical_card(slide, context, item, card_left, card_top, card_width, card_height)
            else:
                self._render_horizontal_card(slide, context, item, card_left, card_top, card_width, card_height)
        
        num_rows = (len(items) + columns - 1) // columns
        return Emu(num_rows * (int(card_height) + int(self.CARD_MARGIN)))
    
    def get_required_height(
        self, 
        context: RenderContext, 
        data: Dict[str, Any],
        **kwargs
    ) -> Emu:
        """필요한 높이 계산"""
        items = data.get('items', [])
        columns = data.get('columns', 2)
        style = data.get('style', 'horizontal')
        
        card_height = self.CARD_HEIGHT_VERTICAL if style == 'vertical' else self.CARD_HEIGHT_HORIZONTAL
        
        if not items:
            return Emu(0)
        
        num_rows = (len(items) + columns - 1) // columns
        return Emu(num_rows * (int(card_height) + int(self.CARD_MARGIN)))
    
    def _render_horizontal_card(
        self,
        slide: Slide,
        context: RenderContext,
        item: Dict[str, Any],
        left: Emu,
        top: Emu,
        width: Emu,
        height: Emu
    ) -> None:
        """가로 레이아웃 카드 렌더링"""
        # 카드 배경
        card_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, height
        )
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = context.theme.colors.background
        card_bg.line.color.rgb = context.theme.colors.secondary
        card_bg.line.width = Pt(1)
        
        # 번호 영역 (왼쪽)
        number = item.get('number', '01')
        number_width = Emu(Cm(1.5))
        number_height = height
        
        number_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left, top, number_width, height
        )
        number_bg.fill.solid()
        number_bg.fill.fore_color.rgb = context.theme.colors.primary
        number_bg.line.fill.background()
        
        number_shape = slide.shapes.add_textbox(left, top, number_width, height)
        tf = number_shape.text_frame
        tf.auto_size = None
        p = tf.paragraphs[0]
        p.text = number
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.title_font
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.text_on_dark
        
        # 텍스트 영역 (오른쪽)
        text_left = Emu(int(left) + int(number_width) + Emu(Cm(0.3)))
        text_width = Emu(int(width) - int(number_width) - Emu(Cm(0.5)))
        
        # 제목
        title = item.get('title', '')
        title_top = Emu(int(top) + Emu(Cm(0.3)))
        title_height = Emu(Cm(0.8))
        
        title_shape = slide.shapes.add_textbox(text_left, title_top, text_width, title_height)
        tf = title_shape.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.text_primary
        
        # 설명
        description = item.get('description', '')
        if description:
            desc_top = Emu(int(top) + Emu(Cm(1.2)))
            desc_height = Emu(Cm(1))
            
            desc_shape = slide.shapes.add_textbox(text_left, desc_top, text_width, desc_height)
            tf = desc_shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = description
            p.alignment = PP_ALIGN.LEFT
            
            run = p.runs[0] if p.runs else p.add_run()
            run.font.name = context.theme.fonts.caption_font
            run.font.size = Pt(10)
            run.font.color.rgb = context.theme.colors.text_secondary
    
    def _render_vertical_card(
        self,
        slide: Slide,
        context: RenderContext,
        item: Dict[str, Any],
        left: Emu,
        top: Emu,
        width: Emu,
        height: Emu
    ) -> None:
        """세로 레이아웃 카드 렌더링"""
        # 카드 배경
        card_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, height
        )
        card_bg.fill.solid()
        card_bg.fill.fore_color.rgb = context.theme.colors.background
        card_bg.line.color.rgb = context.theme.colors.secondary
        card_bg.line.width = Pt(1)
        
        # 번호 (상단 중앙)
        number = item.get('number', '01')
        number_top = Emu(int(top) + Emu(Cm(0.3)))
        number_height = Emu(Cm(1))
        
        number_shape = slide.shapes.add_textbox(left, number_top, width, number_height)
        tf = number_shape.text_frame
        p = tf.paragraphs[0]
        p.text = number
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.title_font
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.accent
        
        # 제목
        title = item.get('title', '')
        title_top = Emu(int(top) + Emu(Cm(1.5)))
        title_height = Emu(Cm(0.8))
        
        title_shape = slide.shapes.add_textbox(left, title_top, width, title_height)
        tf = title_shape.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.primary
        
        # 설명
        description = item.get('description', '')
        if description:
            desc_top = Emu(int(top) + Emu(Cm(2.5)))
            desc_height = Emu(Cm(1.3))
            
            desc_shape = slide.shapes.add_textbox(
                Emu(int(left) + Emu(Cm(0.2))),
                desc_top,
                Emu(int(width) - Emu(Cm(0.4))),
                desc_height
            )
            tf = desc_shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = description
            p.alignment = PP_ALIGN.CENTER
            
            run = p.runs[0] if p.runs else p.add_run()
            run.font.name = context.theme.fonts.caption_font
            run.font.size = Pt(10)
            run.font.color.rgb = context.theme.colors.text_secondary


# CardGridComponent는 다른 카드 컴포넌트들을 조합하여 사용하는 편의 클래스
class CardGridComponent(BaseComponent):
    """카드 그리드 컴포넌트
    
    다양한 카드 타입을 그리드로 배치합니다.
    내부적으로 적절한 카드 컴포넌트를 선택하여 사용합니다.
    """
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any],
        top: Emu,
        **kwargs
    ) -> Emu:
        """카드 그리드 렌더링"""
        card_type = data.get('card_type', 'icon')
        
        component_map = {
            'big_number': BigNumberCardComponent(),
            'icon': IconCardComponent(),
            'feature': FeatureCardComponent(),
        }
        
        component = component_map.get(card_type, IconCardComponent())
        return component.render(slide, context, data, top, **kwargs)
    
    def get_required_height(
        self, 
        context: RenderContext, 
        data: Dict[str, Any],
        **kwargs
    ) -> Emu:
        """필요한 높이 계산"""
        card_type = data.get('card_type', 'icon')
        
        component_map = {
            'big_number': BigNumberCardComponent(),
            'icon': IconCardComponent(),
            'feature': FeatureCardComponent(),
        }
        
        component = component_map.get(card_type, IconCardComponent())
        return component.get_required_height(context, data, **kwargs)
