"""
본문 슬라이드 컴포넌트

PPT 본문 슬라이드를 생성합니다.
동국제강 템플릿 Layout 3-5 기반입니다.

레이아웃 옵션:
    - Layout 3: Action Title 사용 + 본문
    - Layout 4: Action Title 사용 + 자유 콘텐츠 영역
    - Layout 5: Action Title 미사용 + 넓은 본문
"""

from typing import Dict, Any, List, Optional
from pptx.slide import Slide
from pptx.util import Emu, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from ..base import BaseSlideComponent, RenderContext
from ..factory import register_component, VisualComponentFactory


@register_component('content', category='slide')
class ContentSlideComponent(BaseSlideComponent):
    """본문 슬라이드 컴포넌트
    
    템플릿 Layout 3-5 기반:
        - Main Title (19pt, 본고딕 Medium, 네이비)
        - 구분선
        - Action Title (17pt, 선택적)
        - 본문 영역 (시각 요소 컴포넌트 조합)
        - 바닥글, 슬라이드 번호
        
    Data Schema:
        {
            "main_title": "Main Title 영역",
            "action_title": "Action Title (선택적)",  
            "use_action_title": true,  # Action Title 사용 여부
            "body": [...],  # 불릿 포인트 리스트
            "visual_elements": [  # 시각 요소 리스트
                {"type": "table", "data": {...}},
                {"type": "big_numbers", "data": {...}},
                ...
            ],
            "slide_number": 3,
            "footer_text": "작성자_00000팀"
        }
    """
    
    # 레이아웃 상수 (PPT기본양식 분석 결과 기반)
    # Main Title 영역
    MAIN_TITLE_LEFT = Emu(272833)
    MAIN_TITLE_TOP = Emu(171278)
    MAIN_TITLE_WIDTH = Emu(6172693)
    MAIN_TITLE_HEIGHT = Emu(369332)
    
    # 구분선
    SEPARATOR_LEFT = Emu(270064)
    SEPARATOR_TOP = Emu(540000)
    SEPARATOR_WIDTH = Emu(9362886)
    SEPARATOR_HEIGHT = Emu(6350)  # 약 0.5pt
    
    # Action Title 영역 (Layout 3, 4용)
    ACTION_TITLE_LEFT = Emu(272724)
    ACTION_TITLE_TOP = Emu(576000)
    ACTION_TITLE_WIDTH = Emu(9355229)
    ACTION_TITLE_HEIGHT = Emu(708943)
    
    # 본문 영역 (Action Title 있을 때)
    BODY_WITH_ACTION_TOP = Emu(1431130)
    BODY_WITH_ACTION_HEIGHT = Emu(877888)
    
    # 본문 영역 (Action Title 없을 때, Layout 5)
    BODY_NO_ACTION_TOP = Emu(931201)
    BODY_NO_ACTION_HEIGHT = Emu(877888)
    
    # 콘텐츠 영역 (시각 요소용)
    CONTENT_LEFT = Emu(270064)
    CONTENT_WIDTH = Emu(9360550)
    CONTENT_TOP_WITH_ACTION = Emu(1500000)  # Action Title 아래
    CONTENT_TOP_NO_ACTION = Emu(600000)  # 구분선 바로 아래
    
    # 로고 위치
    LOGO_LEFT = Emu(270064)
    LOGO_TOP = Emu(6579886)
    LOGO_WIDTH = Emu(1215734)
    LOGO_HEIGHT = Emu(181025)
    
    # 글머리 기호
    BULLET_CHAR_L1 = "▐"
    BULLET_CHAR_L2 = "•"
    BULLET_CHAR_L3 = "–"
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any]
    ) -> None:
        """본문 슬라이드 렌더링
        
        Args:
            slide: 렌더링할 슬라이드 객체
            context: 렌더링 컨텍스트
            data: 슬라이드 데이터
        """
        use_action_title = data.get('use_action_title', True)
        
        # 1. 상단 그라데이션 바
        self._add_header_bar(slide, context)
        
        # 2. Main Title
        main_title = data.get('main_title', data.get('title', ''))
        self._add_main_title(slide, context, main_title)
        
        # 3. 구분선
        self._add_separator(slide, context)
        
        # 4. Action Title (선택적)
        if use_action_title:
            action_title = data.get('action_title', '')
            if action_title:
                self._add_action_title(slide, context, action_title)
        
        # 5. 콘텐츠 영역 (본문 또는 시각 요소)
        content_top = self.CONTENT_TOP_WITH_ACTION if use_action_title else self.CONTENT_TOP_NO_ACTION
        
        # 불릿 포인트 본문
        body = data.get('body', [])
        if body:
            self._add_body_bullets(slide, context, body, content_top)
            content_top = Emu(int(content_top) + int(self.BODY_WITH_ACTION_HEIGHT))
        
        # 시각 요소들
        visual_elements = data.get('visual_elements', [])
        self._render_visual_elements(slide, context, visual_elements, content_top)
        
        # 6. 바닥글 및 슬라이드 번호
        slide_number = data.get('slide_number')
        footer_text = data.get('footer_text', '')
        self._add_footer(slide, context, slide_number, footer_text)
        
        # 7. 로고
        if context.theme.logo_path:
            self._add_logo_internal(slide, context)
    
    def _add_main_title(
        self,
        slide: Slide,
        context: RenderContext,
        title: str
    ) -> None:
        """Main Title 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            title: 제목 텍스트
        """
        shape = slide.shapes.add_textbox(
            self.MAIN_TITLE_LEFT,
            self.MAIN_TITLE_TOP,
            self.MAIN_TITLE_WIDTH,
            self.MAIN_TITLE_HEIGHT
        )
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(context.theme.fonts.main_title_size)
        run.font.color.rgb = context.theme.colors.primary
    
    def _add_separator(
        self,
        slide: Slide,
        context: RenderContext
    ) -> None:
        """구분선 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
        """
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            self.SEPARATOR_LEFT,
            self.SEPARATOR_TOP,
            self.SEPARATOR_WIDTH,
            self.SEPARATOR_HEIGHT
        )
        
        # 연회색 (bg1 65%)
        shape.fill.solid()
        shape.fill.fore_color.rgb = context.theme.colors.secondary
        shape.line.fill.background()
    
    def _add_action_title(
        self,
        slide: Slide,
        context: RenderContext,
        title: str
    ) -> None:
        """Action Title 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            title: Action Title 텍스트
        """
        shape = slide.shapes.add_textbox(
            self.ACTION_TITLE_LEFT,
            self.ACTION_TITLE_TOP,
            self.ACTION_TITLE_WIDTH,
            self.ACTION_TITLE_HEIGHT
        )
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(context.theme.fonts.action_title_size)
        run.font.color.rgb = context.theme.colors.text_primary
    
    def _add_body_bullets(
        self,
        slide: Slide,
        context: RenderContext,
        body: List[Dict[str, Any]],
        top: Emu
    ) -> Emu:
        """본문 불릿 포인트 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            body: 불릿 포인트 리스트
            top: 시작 Y 위치
            
        Returns:
            사용된 높이
        """
        shape = slide.shapes.add_textbox(
            self.CONTENT_LEFT,
            top,
            self.CONTENT_WIDTH,
            self.BODY_WITH_ACTION_HEIGHT
        )
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        for i, item in enumerate(body):
            if isinstance(item, str):
                text = item
                level = 0
            else:
                text = item.get('text', '')
                level = item.get('level', 0)
            
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            # 글머리 기호 선택
            bullet_char = self._get_bullet_char(level)
            indent = self._get_indent(level)
            
            p.text = f"{bullet_char} {text}"
            p.alignment = PP_ALIGN.LEFT
            p.level = min(level, 4)
            
            # 폰트 설정
            for run in p.runs:
                run.font.name = context.theme.fonts.body_font
                run.font.size = Pt(context.theme.fonts.body_size)
                run.font.color.rgb = context.theme.colors.text_primary
        
        return self.BODY_WITH_ACTION_HEIGHT
    
    def _get_bullet_char(self, level: int) -> str:
        """레벨에 따른 글머리 기호 반환"""
        if level == 0:
            return self.BULLET_CHAR_L1
        elif level == 1:
            return self.BULLET_CHAR_L2
        else:
            return self.BULLET_CHAR_L3
    
    def _get_indent(self, level: int) -> Emu:
        """레벨에 따른 들여쓰기 반환"""
        indent_map = {
            0: Emu(228600),
            1: Emu(685800),
            2: Emu(1143000),
            3: Emu(1600200),
            4: Emu(2057400),
        }
        return indent_map.get(level, Emu(228600))
    
    def _render_visual_elements(
        self,
        slide: Slide,
        context: RenderContext,
        elements: List[Dict[str, Any]],
        top: Emu
    ) -> Emu:
        """시각 요소들 렌더링
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            elements: 시각 요소 리스트
            top: 시작 Y 위치
            
        Returns:
            총 사용된 높이
        """
        current_top = top
        
        for element in elements:
            element_type = element.get('type', '')
            element_data = element.get('data', element)
            
            if not VisualComponentFactory.has_type(element_type):
                continue
            
            try:
                component = VisualComponentFactory.create(element_type)
                height_used = component.render(slide, context, element_data, current_top)
                current_top = Emu(int(current_top) + int(height_used) + Emu(Cm(0.3)))
            except Exception as e:
                # 렌더링 실패 시 건너뛰기
                print(f"Warning: Failed to render {element_type}: {e}")
                continue
        
        return Emu(int(current_top) - int(top))
    
    def _add_logo_internal(
        self,
        slide: Slide,
        context: RenderContext
    ) -> None:
        """로고 이미지 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
        """
        import os
        
        if not context.theme.logo_path or not os.path.exists(context.theme.logo_path):
            return
        
        slide.shapes.add_picture(
            context.theme.logo_path,
            self.LOGO_LEFT,
            self.LOGO_TOP,
            self.LOGO_WIDTH,
            self.LOGO_HEIGHT
        )
