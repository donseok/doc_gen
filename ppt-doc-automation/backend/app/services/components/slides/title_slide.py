"""
표지 슬라이드 컴포넌트

PPT 표지 슬라이드를 생성합니다.
동국제강 템플릿 Layout 1 (White_Big K 버전) 기반입니다.

레이아웃:
    - 상단 그라데이션 바 (빨강→네이비)
    - 문서 제목 (32pt, 본고딕 Bold, 네이비)
    - 부제목/작성자 정보 (14pt, 본고딕 Medium)
    - 하단 로고 이미지
"""

from typing import Dict, Any, Optional
from pptx.slide import Slide
from pptx.util import Emu, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from ..base import BaseSlideComponent, RenderContext
from ..factory import register_component


@register_component('title', category='slide')
class TitleSlideComponent(BaseSlideComponent):
    """표지 슬라이드 컴포넌트
    
    템플릿 Layout 1 (White_Big K 버전) 기반:
        - 상단 그라데이션 바 (빨강→네이비)
        - 문서 제목 (ID:15)
        - 부제목/작성자 정보 (ID:27)
        - 로고 이미지
        
    Data Schema:
        {
            "title": "문서 제목",
            "subtitle": "부제목 또는 작성자 정보",
            "date": "2025-01-03",
            "author": "작성자명"
        }
    """
    
    # 레이아웃 상수 (PPT기본양식 분석 결과 기반)
    TITLE_LEFT = Emu(442133)
    TITLE_TOP = Emu(2276475)
    TITLE_WIDTH = Emu(5398943)
    TITLE_HEIGHT = Emu(956669)
    
    SUBTITLE_LEFT = Emu(442831)
    SUBTITLE_TOP = Emu(3673264)
    SUBTITLE_WIDTH = Emu(5364061)
    SUBTITLE_HEIGHT = Emu(248727)
    
    LOGO_LEFT = Emu(4257434)
    LOGO_TOP = Emu(6473266)
    LOGO_WIDTH = Emu(1391132)
    LOGO_HEIGHT = Emu(207142)
    
    def render(
        self, 
        slide: Slide, 
        context: RenderContext, 
        data: Dict[str, Any]
    ) -> None:
        """표지 슬라이드 렌더링
        
        Args:
            slide: 렌더링할 슬라이드 객체
            context: 렌더링 컨텍스트
            data: 슬라이드 데이터 (title, subtitle, date, author)
        """
        # 1. 상단 그라데이션 바
        self._add_header_bar(slide, context)
        
        # 2. 제목
        title = data.get('title', '')
        self._add_title(slide, context, title)
        
        # 3. 부제목 (작성자/날짜 포함)
        subtitle = self._format_subtitle(data)
        if subtitle:
            self._add_subtitle(slide, context, subtitle)
        
        # 4. 로고 (logo_path가 있는 경우)
        if context.theme.logo_path:
            self._add_logo(slide, context, context.theme.logo_path)
    
    def _add_title(
        self,
        slide: Slide,
        context: RenderContext,
        title: str
    ) -> None:
        """제목 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            title: 제목 텍스트
        """
        shape = slide.shapes.add_textbox(
            self.TITLE_LEFT,
            self.TITLE_TOP,
            self.TITLE_WIDTH,
            self.TITLE_HEIGHT
        )
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        
        # 폰트 설정
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.title_font
        run.font.size = Pt(context.theme.fonts.title_size)
        run.font.bold = True
        run.font.color.rgb = context.theme.colors.primary
    
    def _add_subtitle(
        self,
        slide: Slide,
        context: RenderContext,
        subtitle: str
    ) -> None:
        """부제목 추가
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            subtitle: 부제목 텍스트
        """
        shape = slide.shapes.add_textbox(
            self.SUBTITLE_LEFT,
            self.SUBTITLE_TOP,
            self.SUBTITLE_WIDTH,
            self.SUBTITLE_HEIGHT
        )
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.LEFT
        
        # 폰트 설정
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = context.theme.fonts.body_font
        run.font.size = Pt(context.theme.fonts.subtitle_size)
        run.font.color.rgb = context.theme.colors.text_secondary
    
    def _format_subtitle(self, data: Dict[str, Any]) -> str:
        """부제목 포맷팅
        
        subtitle, author, date 정보를 조합하여 부제목 생성
        
        Args:
            data: 슬라이드 데이터
            
        Returns:
            포맷팅된 부제목 문자열
        """
        parts = []
        
        # 기본 subtitle이 있으면 사용
        if data.get('subtitle'):
            parts.append(data['subtitle'])
        
        # 작성자 정보
        if data.get('author'):
            if data.get('department'):
                parts.append(f"{data['department']} | {data['author']}")
            else:
                parts.append(data['author'])
        
        # 날짜 정보
        if data.get('date'):
            parts.append(data['date'])
        
        return ' | '.join(parts) if parts else ''
    
    def _add_logo(
        self,
        slide: Slide,
        context: RenderContext,
        logo_path: str
    ) -> None:
        """로고 이미지 추가 (표지용 위치)
        
        Args:
            slide: 슬라이드 객체
            context: 렌더링 컨텍스트
            logo_path: 로고 이미지 경로
        """
        import os
        
        if not os.path.exists(logo_path):
            return
        
        slide.shapes.add_picture(
            logo_path,
            self.LOGO_LEFT,
            self.LOGO_TOP,
            self.LOGO_WIDTH,
            self.LOGO_HEIGHT
        )
