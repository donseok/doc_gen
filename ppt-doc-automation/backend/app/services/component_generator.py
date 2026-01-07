"""
컴포넌트 기반 PPT 생성기

JSON v2 형식의 데이터를 컴포넌트들을 조합하여 고품질 PPT로 변환합니다.

Usage:
    from app.services.component_generator import ComponentBasedGenerator
    
    generator = ComponentBasedGenerator(theme="dongkuk")
    prs = generator.generate(json_data)
    prs.save("output.pptx")
"""

import json
import os
import uuid
from typing import Dict, Any, Optional, List
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Emu, Pt, Cm, Inches

from .components.base import RenderContext
from .components.theme import (
    ThemeConfig, DongkukTheme, ModernBlueTheme, 
    CorporateTheme, DarkTheme, get_theme
)
from .components.factory import SlideComponentFactory, VisualComponentFactory

# 슬라이드 컴포넌트 import (팩토리에 자동 등록됨)
from .components.slides import (
    TitleSlideComponent,
    TOCSlideComponent,
    ContentSlideComponent,
    SectionSlideComponent,
)

# 시각 요소 컴포넌트 import (팩토리에 자동 등록됨)
from .components.visual import (
    TableComponent,
    BigNumberCardComponent,
    IconCardComponent,
    FeatureCardComponent,
    CardGridComponent,
    TimelineComponent,
    MilestoneComponent,
    GanttComponent,
)


class ComponentBasedGenerator:
    """컴포넌트 기반 PPT 생성기
    
    JSON v2 형식의 데이터를 입력받아 컴포넌트들을 조합하여 PPT를 생성합니다.
    
    Attributes:
        theme: 적용할 테마
        context: 렌더링 컨텍스트
        
    Usage:
        generator = ComponentBasedGenerator(theme="dongkuk")
        prs = generator.generate(json_data)
        prs.save("output.pptx")
    """
    
    # 슬라이드 크기 상수 (동국제강 템플릿 기준 A4 가로)
    SLIDE_WIDTH = Emu(9906000)   # 약 27.52cm
    SLIDE_HEIGHT = Emu(6858000)  # 약 19.05cm
    
    def __init__(
        self, 
        theme: str = "dongkuk",
        logo_path: Optional[str] = None
    ):
        """생성기 초기화
        
        Args:
            theme: 테마 이름 (dongkuk, modern_blue, corporate, dark)
            logo_path: 로고 이미지 경로 (선택적)
        """
        self.theme_config = get_theme(theme, logo_path)
        self.context = self._create_context()
        self.prs: Optional[Presentation] = None
    
    def _create_context(self) -> RenderContext:
        """렌더링 컨텍스트 생성"""
        return RenderContext(
            theme=self.theme_config,
            slide_width=self.SLIDE_WIDTH,
            slide_height=self.SLIDE_HEIGHT,
            margin_left=Emu(270064),
            margin_right=Emu(270064),
            margin_top=Emu(171278),
            margin_bottom=Emu(250000),
        )
    
    def generate(self, json_data: Dict[str, Any]) -> Presentation:
        """JSON 데이터를 PPT로 변환
        
        Args:
            json_data: 구조화된 프레젠테이션 JSON
            
        Returns:
            Presentation 객체
        """
        # 새 프레젠테이션 생성
        self.prs = Presentation()
        self._setup_slide_size()
        
        # 메타데이터 처리 (선택적)
        metadata = json_data.get('metadata', {})
        
        # 슬라이드 생성
        slides_data = json_data.get('slides', [])
        for i, slide_data in enumerate(slides_data):
            self._render_slide(slide_data, slide_number=i + 1)
        
        return self.prs
    
    def generate_from_file(
        self, 
        json_path: str, 
        output_path: str
    ) -> str:
        """JSON 파일에서 PPT 생성
        
        Args:
            json_path: JSON 파일 경로
            output_path: 출력 PPTX 경로
            
        Returns:
            생성된 파일 경로
        """
        with open(json_path, 'r', encoding='utf-8') as f:
            json_data = json.load(f)
        
        prs = self.generate(json_data)
        prs.save(output_path)
        
        return output_path
    
    def _setup_slide_size(self) -> None:
        """슬라이드 크기 설정"""
        self.prs.slide_width = self.SLIDE_WIDTH
        self.prs.slide_height = self.SLIDE_HEIGHT
    
    def _render_slide(
        self, 
        slide_data: Dict[str, Any],
        slide_number: int = 1
    ) -> None:
        """슬라이드 렌더링
        
        Args:
            slide_data: 슬라이드 데이터
            slide_number: 슬라이드 번호
        """
        slide_type = slide_data.get('slide_type', 'content')
        
        # 빈 레이아웃으로 슬라이드 추가
        blank_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(blank_layout)
        
        # 슬라이드 번호 추가
        slide_data['slide_number'] = slide_number
        
        # 슬라이드 컴포넌트 생성 및 렌더링
        try:
            slide_component = SlideComponentFactory.create(slide_type)
            slide_component.render(slide, self.context, slide_data)
        except ValueError as e:
            print(f"Warning: {e}. Using default content slide.")
            # 기본 콘텐츠 슬라이드로 폴백
            content_component = ContentSlideComponent()
            content_component.render(slide, self.context, slide_data)


class ComponentGeneratorWithTemplate(ComponentBasedGenerator):
    """템플릿 기반 컴포넌트 생성기
    
    기존 PPTX 템플릿을 기반으로 컴포넌트를 렌더링합니다.
    템플릿의 마스터 슬라이드와 레이아웃을 활용합니다.
    """
    
    def __init__(
        self,
        template_path: str,
        theme: str = "dongkuk",
        logo_path: Optional[str] = None
    ):
        """템플릿 기반 생성기 초기화
        
        Args:
            template_path: PPTX 템플릿 경로
            theme: 테마 이름
            logo_path: 로고 이미지 경로
        """
        super().__init__(theme, logo_path)
        self.template_path = template_path
    
    def generate(self, json_data: Dict[str, Any]) -> Presentation:
        """템플릿 기반으로 PPT 생성
        
        Args:
            json_data: 구조화된 프레젠테이션 JSON
            
        Returns:
            Presentation 객체
        """
        # 템플릿에서 프레젠테이션 로드
        if os.path.exists(self.template_path):
            self.prs = Presentation(self.template_path)
            # 기존 슬라이드 모두 제거 (마스터만 유지)
            while len(self.prs.slides) > 0:
                rId = self.prs.slides._sldIdLst[0].rId
                self.prs.part.drop_rel(rId)
                del self.prs.slides._sldIdLst[0]
        else:
            self.prs = Presentation()
            self._setup_slide_size()
        
        # 슬라이드 생성
        slides_data = json_data.get('slides', [])
        for i, slide_data in enumerate(slides_data):
            self._render_slide(slide_data, slide_number=i + 1)
        
        return self.prs


def generate_ppt_from_json(
    json_data: Dict[str, Any],
    output_path: str,
    theme: str = "dongkuk",
    logo_path: Optional[str] = None
) -> str:
    """편의 함수: JSON에서 PPT 생성
    
    Args:
        json_data: 구조화된 프레젠테이션 JSON
        output_path: 출력 PPTX 경로
        theme: 테마 이름
        logo_path: 로고 이미지 경로
        
    Returns:
        생성된 파일 경로
    """
    generator = ComponentBasedGenerator(theme=theme, logo_path=logo_path)
    prs = generator.generate(json_data)
    prs.save(output_path)
    return output_path
