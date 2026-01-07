"""
템플릿 디자인 기반 PPT 생성기

템플릿에서 추출한 디자인 스펙(위치, 크기, 색상, 로고)을 적용하여
새 PPT를 생성합니다.
"""
import json
import os
import uuid
from typing import Dict, List, Any, Optional
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Emu, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def hex_to_rgb(hex_color: str) -> RGBColor:
    """HEX 색상을 RGBColor로 변환"""
    hex_color = hex_color.lstrip("#")
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    return RGBColor(r, g, b)


@dataclass
class DesignSpec:
    """디자인 스펙 데이터 클래스"""
    slide_width: int
    slide_height: int
    layouts: List[Dict]
    logos: List[Dict]
    colors: Dict[str, str]
    fonts: Dict[str, str]

    @classmethod
    def from_json(cls, json_path: str) -> 'DesignSpec':
        """JSON 파일에서 디자인 스펙 로드"""
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        return cls(
            slide_width=data['slide_size']['width_emu'],
            slide_height=data['slide_size']['height_emu'],
            layouts=data['layouts'],
            logos=data['logos'],
            colors=data['colors'],
            fonts=data['fonts']
        )


class TemplateBasedGenerator:
    """
    템플릿 디자인 기반 PPT 생성기

    추출된 디자인 스펙을 적용하여 새 PPT를 생성합니다.
    """

    def __init__(self, design_spec_path: str, logos_dir: str):
        """
        Args:
            design_spec_path: design_spec.json 경로
            logos_dir: 추출된 로고 이미지 디렉토리
        """
        self.spec = DesignSpec.from_json(design_spec_path)
        self.logos_dir = logos_dir
        self.prs = None

        # 색상 설정
        self.colors = {k: hex_to_rgb(v) for k, v in self.spec.colors.items()}

    def generate(self, json_data: Dict) -> Presentation:
        """
        JSON 데이터를 PPT로 변환

        Args:
            json_data: 프레젠테이션 JSON 데이터

        Returns:
            Presentation 객체
        """
        # 빈 프레젠테이션 생성
        self.prs = Presentation()
        self.prs.slide_width = Emu(self.spec.slide_width)
        self.prs.slide_height = Emu(self.spec.slide_height)

        presentation = json_data.get("presentation", {})
        slides = presentation.get("slides", [])

        for slide_data in slides:
            self._add_slide(slide_data)

        return self.prs

    def _add_slide(self, slide_data: Dict):
        """슬라이드 추가"""
        slide_type = slide_data.get("slide_type", "content")
        layout_id = slide_data.get("layout_id", 4)

        # 빈 레이아웃 사용
        blank_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(blank_layout)

        # 슬라이드 타입별 처리
        if slide_type == "title":
            self._render_title_slide(slide, slide_data)
        elif slide_type == "toc":
            self._render_toc_slide(slide, slide_data)
        else:
            self._render_content_slide(slide, slide_data)

    def _get_layout_spec(self, layout_index: int) -> Optional[Dict]:
        """레이아웃 스펙 가져오기"""
        for layout in self.spec.layouts:
            if layout['index'] == layout_index:
                return layout
        return None

    def _add_logo(self, slide, layout_index: int = 0):
        """로고 추가"""
        # 해당 레이아웃의 로고 정보 찾기
        for logo in self.spec.logos:
            if logo['layout_index'] == layout_index:
                logo_path = os.path.join(self.logos_dir, logo['filename'])
                if os.path.exists(logo_path):
                    slide.shapes.add_picture(
                        logo_path,
                        Emu(logo['left_emu']),
                        Emu(logo['top_emu']),
                        Emu(logo['width_emu']),
                        Emu(logo['height_emu'])
                    )
                return

        # 기본 로고 (layout 0)
        for logo in self.spec.logos:
            if logo['layout_index'] == 0:
                logo_path = os.path.join(self.logos_dir, logo['filename'])
                if os.path.exists(logo_path):
                    slide.shapes.add_picture(
                        logo_path,
                        Emu(logo['left_emu']),
                        Emu(logo['top_emu']),
                        Emu(logo['width_emu']),
                        Emu(logo['height_emu'])
                    )
                return

    def _add_header_bar(self, slide):
        """상단 헤더 바 추가"""
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(0), Emu(0),
            Emu(self.spec.slide_width), Emu(120000)  # 약 0.33cm
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.colors['secondary']  # Navy
        bar.line.fill.background()

    def _render_title_slide(self, slide, slide_data: Dict):
        """표지 슬라이드 렌더링"""
        placeholders = slide_data.get("placeholders", {})
        layout_spec = self._get_layout_spec(0)

        # 배경 상단 영역 (네이비)
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(0), Emu(0),
            Emu(self.spec.slide_width), Emu(self.spec.slide_height // 3)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = self.colors['secondary']
        bg_shape.line.fill.background()

        # 제목 (템플릿 스펙에서 위치 가져오기)
        title_spec = None
        if layout_spec:
            for ph in layout_spec['placeholders']:
                if 'TITLE' in ph['type']:
                    title_spec = ph
                    break

        title_left = Emu(title_spec['left_emu']) if title_spec else Cm(1.23)
        title_top = Emu(title_spec['top_emu']) if title_spec else Cm(6.32)
        title_width = Emu(title_spec['width_emu']) if title_spec else Cm(15)
        title_height = Emu(title_spec['height_emu']) if title_spec else Cm(2.66)

        title_box = slide.shapes.add_textbox(
            title_left, title_top, title_width, title_height
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = placeholders.get("title", "")
        p.font.name = self.spec.fonts.get("title", "맑은 고딕")
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = self.colors['text_dark']

        # 부제목
        body_spec = None
        if layout_spec:
            for ph in layout_spec['placeholders']:
                if 'BODY' in ph['type']:
                    body_spec = ph
                    break

        subtitle_left = Emu(body_spec['left_emu']) if body_spec else Cm(1.33)
        subtitle_top = Emu(body_spec['top_emu']) if body_spec else Cm(10.2)
        subtitle_width = Emu(body_spec['width_emu']) if body_spec else Cm(14.9)

        subtitle_box = slide.shapes.add_textbox(
            subtitle_left, subtitle_top, subtitle_width, Cm(1)
        )
        tf = subtitle_box.text_frame
        p = tf.paragraphs[0]
        p.text = placeholders.get("subtitle", "")
        p.font.name = self.spec.fonts.get("body", "맑은 고딕")
        p.font.size = Pt(20)
        p.font.color.rgb = self.colors['gray']

        # 로고
        self._add_logo(slide, 0)

    def _render_toc_slide(self, slide, slide_data: Dict):
        """목차 슬라이드 렌더링"""
        placeholders = slide_data.get("placeholders", {})
        toc_items = placeholders.get("toc_items", [])

        # 헤더 바
        self._add_header_bar(slide)

        # 로고
        self._add_logo(slide, 5)  # 목차 레이아웃

        # 제목
        title_box = slide.shapes.add_textbox(
            Cm(2.44), Cm(4.75), Cm(7.28), Cm(1.33)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = "목차"
        p.font.name = self.spec.fonts.get("title", "맑은 고딕")
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = self.colors['text_dark']

        # 목차 항목들
        start_y = Cm(7.5)
        item_height = Cm(1.3)

        for i, item in enumerate(toc_items[:8]):
            y_pos = start_y + (i * item_height)

            # 번호 원
            num_shape = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Cm(2.5), Emu(int(y_pos)),
                Cm(0.8), Cm(0.8)
            )
            num_shape.fill.solid()
            num_shape.fill.fore_color.rgb = self.colors['secondary']
            num_shape.line.fill.background()

            # 번호 텍스트
            num_tf = num_shape.text_frame
            num_tf.paragraphs[0].text = item.get("number", str(i + 1))
            num_tf.paragraphs[0].font.color.rgb = self.colors['text_light']
            num_tf.paragraphs[0].font.size = Pt(14)
            num_tf.paragraphs[0].font.bold = True
            num_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

            # 제목 텍스트
            title_box = slide.shapes.add_textbox(
                Cm(3.8), Emu(int(y_pos)), Cm(15), Cm(0.8)
            )
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = item.get("title", "")
            p.font.name = self.spec.fonts.get("body", "맑은 고딕")
            p.font.size = Pt(18)
            p.font.color.rgb = self.colors['text_dark']

    def _render_content_slide(self, slide, slide_data: Dict):
        """콘텐츠 슬라이드 렌더링"""
        placeholders = slide_data.get("placeholders", {})
        custom_elements = slide_data.get("custom_elements", [])
        layout_id = slide_data.get("layout_id", 4)

        # 헤더 바
        self._add_header_bar(slide)

        # 로고
        self._add_logo(slide, 1)  # 기본 컨텐츠 레이아웃

        # 메인 제목 (레이아웃 스펙 기반)
        layout_spec = self._get_layout_spec(1)  # 기본 컨텐츠 레이아웃

        title_spec = None
        if layout_spec:
            for ph in layout_spec['placeholders']:
                if 'TITLE' in ph['type']:
                    title_spec = ph
                    break

        # 메인 제목
        main_title = placeholders.get("main_title", "")
        if main_title:
            title_left = Emu(title_spec['left_emu']) if title_spec else Cm(2.48)
            title_top = Cm(1.5)  # 헤더 바 아래

            title_box = slide.shapes.add_textbox(
                title_left, title_top, Cm(20), Cm(1)
            )
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = main_title
            p.font.name = self.spec.fonts.get("title", "맑은 고딕")
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = self.colors['text_dark']

        # 액션 타이틀
        action_title = placeholders.get("action_title", "")
        if action_title:
            action_box = slide.shapes.add_textbox(
                Cm(2.48), Cm(3), Cm(22), Cm(1)
            )
            tf = action_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = action_title
            p.font.name = self.spec.fonts.get("body", "맑은 고딕")
            p.font.size = Pt(14)
            p.font.color.rgb = self.colors['gray']

        # 커스텀 요소 렌더링
        content_top = Cm(5)
        for element in custom_elements:
            self._render_custom_element(slide, element, content_top)

    def _render_custom_element(self, slide, element: Dict, start_top: float):
        """커스텀 시각 요소 렌더링"""
        element_type = element.get("type", "")
        data = element.get("data", {})

        if element_type == "big_number_display":
            self._render_big_numbers(slide, data, start_top)
        elif element_type == "pain_point_cards":
            self._render_pain_cards(slide, data, start_top)
        elif element_type == "icon_card_grid":
            self._render_icon_cards(slide, data, start_top)
        elif element_type == "feature_cards":
            self._render_feature_cards(slide, data, start_top)
        elif element_type == "gantt_timeline":
            self._render_gantt(slide, data, start_top)
        elif element_type == "organization_cards":
            self._render_org_cards(slide, data, start_top)
        elif element_type == "tech_stack_grid":
            self._render_tech_stack(slide, data, start_top)
        elif element_type == "milestone_cards_grid":
            self._render_milestones(slide, data, start_top)
        elif element_type == "two_column_icon_list":
            self._render_two_column_list(slide, data, start_top)
        elif element_type == "risk_matrix_cards":
            self._render_risk_cards(slide, data, start_top)
        elif element_type == "checklist_cards":
            self._render_checklist(slide, data, start_top)
        # 더 많은 요소 타입 추가 가능

    def _render_big_numbers(self, slide, data: Dict, start_top: float):
        """대형 숫자 카드 렌더링"""
        items = data.get("items", [])
        columns = data.get("columns", 4)

        card_width = Cm(5.5)
        card_height = Cm(4)
        margin = Cm(0.5)
        start_left = Cm(2.5)

        for i, item in enumerate(items[:columns]):
            col = i % columns
            left = start_left + col * (card_width + margin)

            # 카드 배경
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(left)), Emu(int(start_top)),
                Emu(int(card_width)), Emu(int(card_height))
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(245, 245, 245)
            card.line.color.rgb = RGBColor(220, 220, 220)

            # 상단 악센트 바
            accent_color = hex_to_rgb(item.get("accent_color", self.spec.colors['secondary']))
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(int(left)), Emu(int(start_top)),
                Emu(int(card_width)), Cm(0.15)
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = accent_color
            accent.line.fill.background()

            # 숫자
            value = item.get("value", "")
            unit = item.get("unit", "")
            num_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.3))), Emu(int(start_top + Cm(0.8))),
                Emu(int(card_width - Cm(0.6))), Cm(1.5)
            )
            tf = num_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"{value}{unit}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = accent_color
            p.alignment = PP_ALIGN.CENTER

            # 레이블
            label = item.get("label", "")
            label_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.3))), Emu(int(start_top + Cm(2.5))),
                Emu(int(card_width - Cm(0.6))), Cm(1)
            )
            tf = label_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = label
            p.font.name = self.spec.fonts.get("body", "맑은 고딕")
            p.font.size = Pt(11)
            p.font.color.rgb = self.colors['gray']
            p.alignment = PP_ALIGN.CENTER

    def _render_pain_cards(self, slide, data: Dict, start_top: float):
        """문제점 카드 렌더링"""
        items = data.get("items", [])
        columns = data.get("columns", 3)

        card_width = Cm(7)
        card_height = Cm(5)
        margin = Cm(0.5)
        start_left = Cm(2)

        for i, item in enumerate(items[:columns]):
            col = i % columns
            left = start_left + col * (card_width + margin)

            # 카드 배경 (빨간색 테두리)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(left)), Emu(int(start_top)),
                Emu(int(card_width)), Emu(int(card_height))
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(255, 245, 245)
            card.line.color.rgb = hex_to_rgb(self.spec.colors['accent'])
            card.line.width = Pt(2)

            # 제목
            title_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.5))), Emu(int(start_top + Cm(0.5))),
                Emu(int(card_width - Cm(1))), Cm(1)
            )
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = item.get("title", "")
            p.font.name = self.spec.fonts.get("title", "맑은 고딕")
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb(self.spec.colors['accent'])

            # 설명
            desc_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.5))), Emu(int(start_top + Cm(1.8))),
                Emu(int(card_width - Cm(1))), Cm(2.5)
            )
            tf = desc_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = item.get("description", "")
            p.font.name = self.spec.fonts.get("body", "맑은 고딕")
            p.font.size = Pt(11)
            p.font.color.rgb = self.colors['text_dark']

    def _render_icon_cards(self, slide, data: Dict, start_top: float):
        """아이콘 카드 그리드"""
        items = data.get("items", [])
        columns = data.get("columns", 4)

        card_width = Cm(5)
        card_height = Cm(4.5)
        margin = Cm(0.5)
        start_left = Cm(2.5)

        for i, item in enumerate(items[:8]):
            col = i % columns
            row = i // columns
            left = start_left + col * (card_width + margin)
            top = start_top + row * (card_height + margin)

            # 카드 배경
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(left)), Emu(int(top)),
                Emu(int(card_width)), Emu(int(card_height))
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(250, 250, 250)
            card.line.color.rgb = RGBColor(230, 230, 230)

            # 제목
            title_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.3))), Emu(int(top + Cm(0.5))),
                Emu(int(card_width - Cm(0.6))), Cm(1)
            )
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = item.get("title", "")
            p.font.name = self.spec.fonts.get("title", "맑은 고딕")
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = self.colors['secondary']

            # 설명
            desc_box = slide.shapes.add_textbox(
                Emu(int(left + Cm(0.3))), Emu(int(top + Cm(1.5))),
                Emu(int(card_width - Cm(0.6))), Cm(2.5)
            )
            tf = desc_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = item.get("description", "")
            p.font.name = self.spec.fonts.get("body", "맑은 고딕")
            p.font.size = Pt(10)
            p.font.color.rgb = self.colors['text_dark']

    def _render_feature_cards(self, slide, data: Dict, start_top: float):
        """기능 카드 렌더링"""
        self._render_icon_cards(slide, data, start_top)

    def _render_gantt(self, slide, data: Dict, start_top: float):
        """간트 차트 렌더링 (간소화)"""
        phases = data.get("phases", [])

        chart_width = Cm(22)
        bar_height = Cm(0.8)
        margin = Cm(0.3)
        start_left = Cm(2.5)

        for i, phase in enumerate(phases[:6]):
            top = start_top + i * (bar_height + margin)

            # 단계명
            label_box = slide.shapes.add_textbox(
                Emu(int(start_left)), Emu(int(top)),
                Cm(4), Emu(int(bar_height))
            )
            tf = label_box.text_frame
            p = tf.paragraphs[0]
            p.text = phase.get("name", "")
            p.font.size = Pt(10)
            p.font.color.rgb = self.colors['text_dark']

            # 바
            bar_left = start_left + Cm(4.5)
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(bar_left)), Emu(int(top)),
                Cm(8), Emu(int(bar_height))
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = hex_to_rgb(phase.get("color", self.spec.colors['secondary']))
            bar.line.fill.background()

    def _render_org_cards(self, slide, data: Dict, start_top: float):
        """조직 카드 렌더링"""
        items = data.get("items", data.get("members", []))
        self._render_icon_cards(slide, {"items": items, "columns": 4}, start_top)

    def _render_tech_stack(self, slide, data: Dict, start_top: float):
        """기술 스택 렌더링"""
        layers = data.get("layers", [])

        layer_width = Cm(22)
        layer_height = Cm(1.5)
        margin = Cm(0.3)
        start_left = Cm(2.5)

        for i, layer in enumerate(layers[:5]):
            top = start_top + i * (layer_height + margin)

            # 레이어 바
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(int(start_left)), Emu(int(top)),
                Emu(int(layer_width)), Emu(int(layer_height))
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = hex_to_rgb(layer.get("color", self.spec.colors['secondary']))
            bar.line.fill.background()

            # 텍스트
            text_box = slide.shapes.add_textbox(
                Emu(int(start_left + Cm(0.5))), Emu(int(top + Cm(0.3))),
                Emu(int(layer_width - Cm(1))), Cm(1)
            )
            tf = text_box.text_frame
            p = tf.paragraphs[0]

            layer_name = layer.get("name", "")
            techs = ", ".join(layer.get("technologies", []))
            p.text = f"{layer_name}: {techs}" if techs else layer_name
            p.font.size = Pt(11)
            p.font.color.rgb = self.colors['text_light']
            p.font.bold = True

    def _render_milestones(self, slide, data: Dict, start_top: float):
        """마일스톤 렌더링"""
        items = data.get("items", [])
        self._render_icon_cards(slide, {"items": items, "columns": 4}, start_top)

    def _render_two_column_list(self, slide, data: Dict, start_top: float):
        """2열 리스트 렌더링"""
        left_items = data.get("left_column", {}).get("items", [])
        right_items = data.get("right_column", {}).get("items", [])

        col_width = Cm(10)
        start_left = Cm(2.5)

        # 왼쪽 열
        for i, item in enumerate(left_items[:5]):
            top = start_top + i * Cm(1.2)
            text_box = slide.shapes.add_textbox(
                Emu(int(start_left)), Emu(int(top)),
                Emu(int(col_width)), Cm(1)
            )
            tf = text_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"• {item.get('text', item) if isinstance(item, dict) else item}"
            p.font.size = Pt(12)
            p.font.color.rgb = self.colors['text_dark']

        # 오른쪽 열
        for i, item in enumerate(right_items[:5]):
            top = start_top + i * Cm(1.2)
            text_box = slide.shapes.add_textbox(
                Emu(int(start_left + col_width + Cm(1))), Emu(int(top)),
                Emu(int(col_width)), Cm(1)
            )
            tf = text_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"• {item.get('text', item) if isinstance(item, dict) else item}"
            p.font.size = Pt(12)
            p.font.color.rgb = self.colors['text_dark']

    def _render_risk_cards(self, slide, data: Dict, start_top: float):
        """위험 카드 렌더링"""
        items = data.get("items", [])
        self._render_pain_cards(slide, {"items": items, "columns": 3}, start_top)

    def _render_checklist(self, slide, data: Dict, start_top: float):
        """체크리스트 렌더링"""
        items = data.get("items", [])

        start_left = Cm(2.5)

        for i, item in enumerate(items[:8]):
            top = start_top + i * Cm(1.2)

            # 체크 박스
            check = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(int(start_left)), Emu(int(top)),
                Cm(0.4), Cm(0.4)
            )
            check.fill.solid()
            check.fill.fore_color.rgb = self.colors['secondary']
            check.line.fill.background()

            # 텍스트
            text_box = slide.shapes.add_textbox(
                Emu(int(start_left + Cm(0.7))), Emu(int(top)),
                Cm(20), Cm(1)
            )
            tf = text_box.text_frame
            p = tf.paragraphs[0]
            text = item.get("title", item) if isinstance(item, dict) else str(item)
            p.text = text
            p.font.size = Pt(12)
            p.font.color.rgb = self.colors['text_dark']


def generate_with_template_design(
    json_path: str,
    design_spec_path: str,
    logos_dir: str,
    output_dir: str
) -> str:
    """
    템플릿 디자인을 적용하여 PPT 생성

    Args:
        json_path: 프레젠테이션 JSON 파일 경로
        design_spec_path: 디자인 스펙 JSON 경로
        logos_dir: 로고 이미지 디렉토리
        output_dir: 출력 디렉토리

    Returns:
        생성된 PPTX 파일 경로
    """
    # JSON 로드
    with open(json_path, 'r', encoding='utf-8') as f:
        json_data = json.load(f)

    # 생성기 초기화 및 생성
    generator = TemplateBasedGenerator(design_spec_path, logos_dir)
    prs = generator.generate(json_data)

    # 파일명 생성
    presentation = json_data.get("presentation", {})
    title = presentation.get("title", "presentation")
    safe_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()[:50]

    output_filename = f"{safe_title}_designed_{uuid.uuid4().hex[:8]}.pptx"
    output_path = os.path.join(output_dir, output_filename)

    # 저장
    prs.save(output_path)

    return output_path
