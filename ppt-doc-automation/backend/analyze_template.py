"""
그룹 공통 PPT템플릿.pptx 분석 스크립트
"""
from pptx import Presentation
from pptx.util import Emu
import os

template_path = r"d:\doc_gen\ppt-doc-automation\templates\pptx\group\그룹 공통 PPT템플릿.pptx"

print(f"템플릿 파일: {template_path}")
print(f"파일 존재: {os.path.exists(template_path)}")
print()

prs = Presentation(template_path)

print(f"=== 기본 정보 ===")
print(f"슬라이드 너비: {prs.slide_width} EMU ({prs.slide_width / 914400:.2f} 인치)")
print(f"슬라이드 높이: {prs.slide_height} EMU ({prs.slide_height / 914400:.2f} 인치)")
print(f"총 슬라이드 수: {len(prs.slides)}")
print(f"총 레이아웃 수: {len(prs.slide_layouts)}")
print()

print(f"=== 레이아웃 목록 ===")
for i, layout in enumerate(prs.slide_layouts):
    print(f"  [{i}] {layout.name}")
print()

print(f"=== 각 슬라이드 상세 ===")
for i, slide in enumerate(prs.slides):
    print(f"\n--- 슬라이드 {i+1} ---")
    print(f"  레이아웃: {slide.slide_layout.name}")
    print(f"  도형 개수: {len(slide.shapes)}")
    
    for shape in slide.shapes:
        shape_info = f"    - [{shape.shape_id}] {shape.name}"
        if hasattr(shape, 'text') and shape.text:
            text_preview = shape.text[:50].replace('\n', ' ')
            shape_info += f" | 텍스트: '{text_preview}...'" if len(shape.text) > 50 else f" | 텍스트: '{text_preview}'"
        if hasattr(shape, 'placeholder_format') and shape.placeholder_format:
            pf = shape.placeholder_format
            shape_info += f" | placeholder_idx={pf.idx}"
        print(shape_info)

print("\n=== 분석 완료 ===")
